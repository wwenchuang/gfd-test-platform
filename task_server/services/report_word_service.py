"""Generate self-contained Word reports from the platform's report Markdown."""

import io
import re


def render_report_docx(markdown):
    # Existing platform dependency; lazy import keeps other downloads independent.
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number'):
        style = document.styles[name]
        style.font.name = 'Arial'
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn(f'w:{attr}'), None)
        fonts.set(qn('w:eastAsia'), '宋体')
        style.font.color.rgb = RGBColor.from_string('172033')
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for border in document.styles['Title'].element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
    document.styles['Title'].font.size = Pt(22)
    document.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    document.styles['Heading 1'].font.size = Pt(15)
    document.styles['Heading 2'].font.size = Pt(12)

    def safe_text(value):
        # XML 1.0 rejects C0 controls; retain tabs/newlines and all normal Unicode.
        return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)

    def add_table(rows):
        count = max(map(len, rows))
        table = document.add_table(rows=0, cols=count)
        table.style = 'Table Grid'
        table.autofit = False
        weights = [min(32, max(5, max(len(row[i]) if i < len(row) else 0 for row in rows))) for i in range(count)]
        widths = [Inches(7 * weight / sum(weights)) for weight in weights]
        for column, width in zip(table.columns, widths):
            column.width = width
        for index, values in enumerate(rows):
            row = table.add_row()
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]
                cell.text = safe_text(values[i] if i < len(values) else '')
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell in row.cells:
                    shade = OxmlElement('w:shd')
                    shade.set(qn('w:fill'), 'EEF4FB')
                    cell._tc.get_or_add_tcPr().append(shade)
        document.add_paragraph().paragraph_format.space_after = Pt(0)

    pending_rows = []
    for raw in str(markdown or '').splitlines():
        line = raw.strip()
        if line.startswith('|') and line.endswith('|'):
            cells = [cell.strip().replace(r'\|', '|') for cell in re.split(r'(?<!\\)\|', line[1:-1])]
            if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                pending_rows.append(cells)
            continue
        if pending_rows:
            add_table(pending_rows)
            pending_rows = []
        if not line:
            continue
        heading = re.match(r'^(#{1,3})\s+(.+)', line)
        if heading:
            level = len(heading[1])
            document.add_paragraph(safe_text(heading[2]), 'Title' if level == 1 else f'Heading {level - 1}')
        elif re.match(r'^\d+\.\s+', line):
            paragraph = document.add_paragraph(safe_text(line))
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        elif line.startswith('- '):
            document.add_paragraph(safe_text(line[2:]), 'List Bullet')
        else:
            document.add_paragraph(safe_text(line))
    if pending_rows:
        add_table(pending_rows)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
