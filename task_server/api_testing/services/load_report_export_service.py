"""Leadership document and spreadsheet from one authorized report snapshot."""
from io import BytesIO
import json
import math


def display(value):
    if value is None or (isinstance(value, float) and not math.isfinite(value)):
        return '未采集'
    if isinstance(value, bool):
        return '是' if value else '否'
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def workload_label(workload):
    workload = workload or {}
    names = {'constant-vus': '固定并发', 'constant-arrival-rate': '固定到达率', 'ramping-vus': '分阶段并发', 'ramping-arrival-rate': '分阶段到达率'}
    fields = {'vus': '并发用户 VU', 'duration_seconds': '持续秒数', 'max_vus': '最多 VU', 'pre_allocated_vus': '预分配 VU', 'start_vus': '起始 VU'}
    parts = [names.get(workload.get('executor'), '历史记录未填写模型')]
    parts.extend(f'{label}：{display(workload[key])}' for key, label in fields.items() if key in workload)
    time_unit = workload.get('time_unit') or '1s'
    unit_label = {'1s': '1 秒', '1m': '1 分钟'}.get(time_unit, str(time_unit))
    for key, title in (('rate', '开始迭代速率'), ('start_rate', '起始迭代速率')):
        if key in workload:
            parts.append(f'{title}：{display(workload[key])} 次 / {unit_label}')
    stage_unit = f'次 / {unit_label}' if 'arrival-rate' in str(workload.get('executor')) else 'VU'
    parts.extend(f'阶段 {i + 1}：{stage.get("duration_seconds")} 秒内到 {stage.get("target")} {stage_unit}' for i, stage in enumerate(workload.get('stages') or []))
    return '；'.join(parts)


def stop_label(policy):
    if not policy: return '未配置自动停止'
    return f'每节点累计 HTTP 错误率超过 {policy["http_error_rate"] * 100:g}%，观察宽限 {policy["grace_seconds"]} 秒；节点失败后通知其他节点停止。不会取消服务端已受理的任务。'


def percentage(value):
    return f'{value * 100:.2f}%' if isinstance(value, (int, float)) and math.isfinite(value) else None


def generator_rows(report):
    rows = []
    for node in report.get('agents', []):
        resource = node.get('load_generator_resources') or {}
        for point in resource.get('samples') or []:
            rows.append([node.get('name'), point.get('sampled_at'), point.get('cpu_scope'), point.get('cpu_used_cores'), point.get('cpu_limit_cores'), point.get('cpu_percent'), point.get('cpu_limit_source'), point.get('memory_scope'), point.get('memory_used_bytes'), point.get('memory_limit_bytes')])
    return rows


def trend_png(points, unit):
    """Bounded static chart. Gaps remain gaps; no interpolation or zero filling."""
    from datetime import datetime, timezone
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.dates import DateFormatter
    values = []
    for timestamp, value in points:
        try: when = datetime.fromtimestamp(timestamp, timezone.utc) if isinstance(timestamp, (int, float)) else datetime.fromisoformat(str(timestamp).replace('Z', '+00:00')).astimezone(timezone.utc)
        except (TypeError, ValueError): continue
        number = value if isinstance(value, (int, float)) and math.isfinite(value) else float('nan')
        values.append((when, number))
    if not any(math.isfinite(value) for _, value in values): return None
    values.sort(key=lambda pair: pair[0])
    gaps = sorted((b[0]-a[0]).total_seconds() for a,b in zip(values,values[1:]) if b[0]>a[0])
    step = gaps[len(gaps)//2] if gaps else 0
    expanded = []
    for index, point in enumerate(values):
        if index and step and (point[0]-values[index-1][0]).total_seconds()>step*1.5: expanded.append((point[0], float('nan')))
        expanded.append(point)
    values=expanded
    fig = Figure(figsize=(7.2, 2.25), dpi=140, layout='constrained')
    FigureCanvasAgg(fig)
    ax = fig.subplots(); ax.plot([x for x, _ in values], [y for _, y in values], color='#186caa', linewidth=1.4, marker='.', markersize=3)
    ax.set_ylabel(unit); ax.set_xlabel('Time (UTC)'); ax.grid(True, color='#dfe7ee', linewidth=.6)
    ax.xaxis.set_major_formatter(DateFormatter('%H:%M:%S'))
    ax.tick_params(labelsize=8)
    result = BytesIO(); fig.savefig(result, format='png'); result.seek(0)
    return result


def label(value):
    return {'host':'整台主机','container':'独立容器','pod':'Pod 内各容器','postgres':'PostgreSQL 指定数据库','completed':'采集完成','partial':'部分缺失','failed':'失败','missing':'采样缺失','not_selected':'未选择监控','collecting':'采集中','host_cpu_cores':'整机逻辑 CPU 核数','host_memory_total_bytes':'主机总内存','container_cpu_quota_cores':'同一容器 CPU quota / period','filesystem_size_bytes':'该文件系统总空间','disk_operations_per_second':'同一设备完成 I/O 速率','not_applicable':'绝对用量或速率，无百分比分母','not_available':'限制未知，不计算百分比','k6_process':'仅 k6 进程','cgroup_v1':'容器/控制组 v1','cgroup_v2':'容器/控制组 v2','visible_cpus':'可见 CPU 核数（非独享配额）','cgroup_quota':'控制组 CPU 配额','unavailable':'未采集'}.get(value, value)


def threshold_rows(report):
    rows=[]
    for item in report.get('thresholds') or []:
        rate=item.get('key','').endswith('_rate')
        expected=percentage(item.get('expected')) if rate else item.get('expected')
        actual=percentage(item.get('actual')) if rate else item.get('actual')
        operator={'less_than':'小于','less_than_or_equal':'不超过','greater_than':'大于','greater_than_or_equal':'不低于'}.get(item.get('operator'),'')
        rows.append([item.get('label') or item.get('key'), f'{operator} {display(expected)}',actual,'缺少样本' if item.get('actual') is None else '通过' if item.get('passed') else '未通过'])
    return rows


def report_sections(report):
    evidence = report.get('evidence') or {}
    context = report.get('test_context') or {}
    transport = report.get('transport') or {}
    count = transport.get('requests') or 0
    latency = report.get('latency') or {}
    from .load_execution_policy import PURPOSES
    sections = [
        ('测试概述', [('执行编号', report.get('run_id')), ('结论', report.get('verdict_label')), ('判断依据', report.get('verdict_explanation')), ('场景', (evidence.get('scenario_snapshot') or {}).get('name')), ('目的', PURPOSES.get(context.get('purpose'), '历史记录未填写')), ('业务版本', context.get('release') or '未记录'), ('数据与账号规模', context.get('data_profile') or '未记录'), ('缓存状态', {'warm':'已预热','cold':'冷缓存','mixed':'混合状态'}.get(context.get('cache_state'), '未记录')), ('测试边界与清理安排', context.get('notes') or '未记录')]),
        ('环境与负载', [('环境', (evidence.get('environment_snapshot') or {}).get('name')), ('压力模型与参数', workload_label(evidence.get('workload_snapshot'))), ('统计时长（秒）', (report.get('statistical_basis') or {}).get('rate_duration_seconds')), ('实际负载判定', (report.get('load_goal') or {}).get('explanation') or (report.get('load_goal') or {}).get('reached')), ('监控证据完整', evidence.get('monitoring_required_complete'))]),
        ('关键结果', [('HTTP请求数', count), ('请求吞吐 RPS', transport.get('requests_per_second') if count else None), ('HTTP错误率', percentage(transport.get('http_error_rate')) if count else None), ('业务断言失败率', percentage((report.get('business') or {}).get('failure_rate')) if (report.get('business') or {}).get('assertions') else None), ('P95 毫秒', latency.get('p95_ms') if count else None), ('P99 毫秒', latency.get('p99_ms') if count else None), ('采样完整性', '计数一致' if (evidence.get('sample_integrity') or {}).get('consistent') else '存在计数差异或未记录；请查看在线报告明细')]),
        ('风险与复验', [('结论适用范围', '仅适用于本次场景、环境、版本、数据和实际达到的负载。短跑不证明最大容量或长期稳定性。'), ('历史可比性', (report.get('comparison') or {}).get('reason') or ('同条件历史可比' if (report.get('comparison') or {}).get('compatible') else '没有可比历史')), ('自动停止策略', stop_label(report.get('stop_policy'))), ('监控状态', label((report.get('monitoring') or {}).get('state')) or '未接入')]),
    ]
    diagnosis = report.get('ai_diagnosis') or {}
    advice = [('诊断状态', '基于同一证据生成，仅作辅助判断' if diagnosis else '尚无匹配当前证据的已完成诊断，不影响基础报告'), ('观察与可能原因', diagnosis.get('conclusion') or '未提供')]
    for index, item in enumerate(diagnosis.get('recommendations') or []):
        advice.extend([(f'建议 {index+1}', item.get('action')), (f'验证方法 {index+1}', item.get('verification'))])
    sections.append(('AI 诊断与建议', advice))
    return sections


def resource_rows(report):
    rows = []
    for service in (report.get('monitoring') or {}).get('services', []):
        for metric in service.get('metrics', []):
            rows.append([service.get('name'), service.get('scope'), service.get('state'), metric.get('label') or metric.get('key'), metric.get('unit'), metric.get('peak'), metric.get('average'), metric.get('denominator'), metric.get('semantics')])
    return rows


def export_report(report, format):
    sections = report_sections(report)
    output = BytesIO()
    if format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.75)
        normal = document.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Songti SC')
        normal.paragraph_format.space_after = Pt(6)
        title = document.add_paragraph('性能测试报告', 'Title')
        for run in title.runs: run.font.color.rgb = RGBColor(0, 0, 0)
        document.add_paragraph(f"结论：{report.get('verdict_label', '证据不足')}。{report.get('verdict_explanation', '')}")
        document.add_paragraph('本文用于评估本次业务负载下的服务表现。未接入、缺失或未执行的指标不作为正常证据。')
        def table(headers, rows):
            tab = document.add_table(rows=1, cols=len(headers)); tab.style = 'Light Shading Accent 1'; tab.autofit = False
            for cell, value in zip(tab.rows[0].cells, headers): cell.text = display(value)
            repeat = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                for cell, value in zip(tab.add_row().cells, values): cell.text = display(value)
            if len(headers) == 2:
                tab.columns[0].width = Inches(1.65); tab.columns[1].width = Inches(5.1)
                for row in tab.rows:
                    row.cells[0].width = Inches(1.65); row.cells[1].width = Inches(5.1)
            for row in tab.rows:
                no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for name, rows in sections:
            document.add_heading(name, level=1); table(['项目', '记录'], rows)
        document.add_heading('性能验收依据', level=1)
        if threshold_rows(report): table(['标准', '要求', '实际', '结果'], threshold_rows(report))
        else: document.add_paragraph('未配置性能验收阈值，请结合原始数据确认业务要求。')
        document.add_heading('响应时间趋势', level=1)
        chart = trend_png([(p.get('started_at'), p.get('p95_ms') if p.get('requests') else None) for p in report.get('series', [])], 'P95 (ms)')
        if chart: document.add_picture(chart, width=Inches(6.5))
        else: document.add_paragraph('没有可绘制的响应时间窗口。')
        document.add_paragraph('纵轴：P95 响应时间（毫秒）。横轴统一 UTC；每点为一个统计窗口，分位数不直接平均。缺失值留空。')
        document.add_heading('接口结果明细', level=1)
        table(['步骤', '请求数', 'HTTP错误率', 'P95 毫秒'], [[s.get('name'), s.get('requests'), percentage(s.get('http_error_rate')) if s.get('requests') else None, s.get('p95_ms') if s.get('requests') else None] for s in report.get('steps', [])])
        document.add_heading('资源监控记录', level=1)
        document.add_paragraph('主机资源不能替代单个服务资源；详细采样数据见同一执行的 Excel 明细和在线曲线。')
        for row in resource_rows(report):
            document.add_heading(f'{display(row[0])} — {display(row[3])}', level=2)
            table(['项目', '记录'], [('范围', label(row[1])), ('状态', label(row[2])), ('峰值', f'{display(row[5])} {display(row[4])}'), ('保留样本平均值', f'{display(row[6])} {display(row[4])}（非服务汇总）'), ('分母', label(row[7])), ('统计含义', row[8])])
        if not resource_rows(report): document.add_paragraph('未接入或没有可用资源样本，不能判断资源是否正常。')
        chart_count = 0
        for service in (report.get('monitoring') or {}).get('services', []):
            for metric in service.get('metrics', []):
                for series in metric.get('series', []):
                    if chart_count >= 8: continue
                    chart = trend_png([(p.get('timestamp'), p.get('value')) for p in series.get('points', [])], str(metric.get('unit') or 'value'))
                    if chart:
                        document.add_heading(f'{service.get("name")} · {metric.get("label") or metric.get("key")}', level=2)
                        document.add_paragraph('实例：' + display(series.get('labels'))).paragraph_format.keep_with_next = True
                        document.add_picture(chart, width=Inches(6.5)); chart_count += 1
        document.add_paragraph('Word 最多展示前 8 条有样本的资源曲线；全部实例、单位及数据点保留在 Excel 和在线报告。')
        document.add_heading('压力机运行资源', level=1)
        document.add_paragraph('下表是发压节点运行时的数据，不是被测服务 CPU。cgroup 包含 Agent 及子进程；k6_process 仅该进程。').paragraph_format.keep_with_next = True
        for node in report.get('agents', []):
            samples = (node.get('load_generator_resources') or {}).get('samples') or []
            def peak(key):
                values = [p.get(key) for p in samples if isinstance(p.get(key), (int, float))]
                return max(values) if values else None
            table(['项目', '记录'], [('节点', node.get('name')), ('保留样本数', len(samples)), ('CPU 峰值（核）', peak('cpu_used_cores')), ('内存峰值（字节）', peak('memory_used_bytes')), ('CPU 范围 / 分母', f'{label(samples[-1].get("cpu_scope"))} / {label(samples[-1].get("cpu_limit_source"))}' if samples else None)])
        if not generator_rows(report): document.add_paragraph('未采集压力机运行资源，不能排除压力机瓶颈。')
        # Declare East Asian fonts on headings and table runs as well as Normal.
        for style in document.styles:
            if style.type == 1:
                style.font.name = 'Arial'
                style.font.color.rgb = RGBColor(0, 0, 0)
                style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        for paragraph in list(document.paragraphs) + [p for t in document.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
            for run in paragraph.runs:
                run.font.name = 'Arial'
                run.font.color.rgb = RGBColor(0, 0, 0)
                run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        footer = document.sections[0].footer.paragraphs[0]
        footer.add_run('性能测试报告 · ' + str(report.get('run_id') or '') + ' · 第 ')
        page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE'); footer._p.append(page)
        footer.add_run(' 页')
        for run in footer.runs:
            run.font.size = Pt(8); run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        document.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        workbook = Workbook(); workbook.remove(workbook.active)
        def sheet(name, headers, rows):
            ws = workbook.create_sheet(name); ws.append(headers)
            for row in rows:
                ws.append([value if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else display(value) for value in row])
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str): cell.data_type = 's'  # never execute spreadsheet formulas
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in ws[1]: cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='164C68')
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for column in ws.columns: ws.column_dimensions[column[0].column_letter].width = 28 if len(headers)>2 else 55
            return ws
        for name, rows in sections: sheet(name, ['项目', '记录'], rows)
        sheet('验收阈值', ['标准','要求','实际','结果'],threshold_rows(report))
        sheet('接口结果', ['步骤ID', '名称', '请求数', 'HTTP错误率', '业务失败率', 'P95毫秒'], [[s.get('id'),s.get('name'),s.get('requests'),s.get('http_error_rate') if s.get('requests') else None,s.get('business_failure_rate') if s.get('requests') else None,s.get('p95_ms') if s.get('requests') else None] for s in report.get('steps',[])])
        sheet('资源摘要', ['目标', '范围', '状态', '指标', '单位', '峰值', '平均值', '分母', '含义'], resource_rows(report))
        points=[]
        for service in (report.get('monitoring') or {}).get('services',[]):
            for metric in service.get('metrics',[]):
                for series in metric.get('series',[]):
                    for point in series.get('points',[]):
                        points.append([service.get('name'), metric.get('key'), series.get('labels'), point.get('timestamp'), point.get('source_timestamp'), point.get('value'), metric.get('unit'), point.get('denominator_value'), point.get('used_value')])
        sheet('资源采样', ['目标','指标','实例标签','查询时间','源时间','数值','单位','分母数值','已用数值'],points)
        sheet('时间趋势', ['时段','请求数','HTTP失败数','P95毫秒'], [[s.get('started_at'),s.get('requests'),s.get('http_failures'),s.get('p95_ms') if s.get('requests') else None] for s in report.get('series',[])])
        sheet('压力机采样', ['节点','采样时间','CPU范围','已用CPU核','CPU分母核','CPU百分比','CPU分母来源','内存范围','已用字节','限制字节'],generator_rows(report))
        workbook.save(output)
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        raise ValueError('仅支持 Word docx 或 Excel xlsx')
    return output.getvalue(), mime
