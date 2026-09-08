from io import BytesIO
from docx import Document
from openpyxl import load_workbook
import pytest
from task_server.api_testing.services.load_report_export_service import export_report
from task_server.api_testing.services.load_report_export_service import workload_label


def test_export_workload_preserves_arrival_rate_time_unit():
    text = workload_label({'executor':'constant-arrival-rate','rate':120,'time_unit':'1m','duration_seconds':60})
    assert '120 次 / 1 分钟' in text
    assert '每秒开始迭代数：120' not in text
    staged = workload_label({'executor':'ramping-arrival-rate','start_rate':30,'time_unit':'1m','stages':[{'duration_seconds':60,'target':120}]})
    assert '30 次 / 1 分钟' in staged and '120 次 / 1 分钟' in staged
    assert '20 VU' in workload_label({'executor':'ramping-vus','stages':[{'duration_seconds':10,'target':20}]})

REPORT={'run_id':'r-test','state':'finished','verdict_label':'证据不足','verdict_explanation':'没有业务资源数据','transport':{'requests':0,'http_error_rate':0},'latency':{'p95_ms':0},'evidence':{'scenario_snapshot':{'name':'=NOT_A_FORMULA()'},'environment_snapshot':{'name':'测试环境'}},'steps':[{'id':'s','name':'中文接口','requests':0,'p95_ms':0}],'test_context':{'purpose':'smoke','notes':'不触发打印'},'monitoring':{'state':'not_selected','services':[]}}

def test_word_has_summary_and_missing_data_not_healthy_zero():
    content,mime=export_report(REPORT,'docx')
    doc=Document(BytesIO(content))
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert '性能测试报告' in text and '证据不足' in text and '未采集' in text
    assert '不触发打印' in text and mime.endswith('document')
    assert 'P95 毫秒\n0' not in text

def test_excel_preserves_numbers_and_never_formula_injection():
    content,_=export_report(REPORT,'xlsx')
    wb=load_workbook(BytesIO(content),data_only=False)
    assert wb['关键结果']['B2'].value == 0
    cells=[c for row in wb['测试概述'] for c in row if c.value=='=NOT_A_FORMULA()']
    assert len(cells)==1 and cells[0].data_type=='s'
    assert wb['接口结果']['F2'].value=='未采集'
    assert '资源采样' in wb.sheetnames

def test_export_rejects_unknown_format():
    with pytest.raises(ValueError): export_report(REPORT,'html')


def test_word_and_excel_include_runtime_and_draw_real_numeric_resource_times():
    report = {**REPORT, 'ai_diagnosis':{'conclusion':'需要核对来源','recommendations':[{'action':'检查资源限制','verification':'同条件复验'}]},
              'monitoring': {'state':'completed','services':[{'name':'主机','scope':'host','state':'completed','metrics':[{'key':'cpu_percent','unit':'%','series':[{'labels':{'instance':'local:9100'},'points':[{'timestamp':1788825600,'value':30},{'timestamp':1788825615,'value':None},{'timestamp':1788825630,'value':50}]}]}]}]},
              'agents':[{'name':'压力机','load_generator_resources':{'samples':[{'sampled_at':'2026-09-08T00:00:00Z','cpu_scope':'k6_process','cpu_limit_source':'visible_cpus','cpu_used_cores':.5,'memory_used_bytes':1048576}]}}]}
    word,_=export_report(report,'docx');doc=Document(BytesIO(word))
    assert len(doc.inline_shapes)==1
    text=' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert '同条件复验' in text and '仅 k6 进程 / 可见 CPU 核数（非独享配额）' in text
    excel,_=export_report(report,'xlsx');wb=load_workbook(BytesIO(excel))
    assert wb['压力机采样']['D2'].value==.5
    assert wb['资源采样']['F3'].value=='未采集'


def test_pod_state_export_label_keeps_scope_distinct_from_container_usage():
    from task_server.api_testing.services.load_report_export_service import label
    assert label('pod_state') == '指定 Pod 状态'

def test_export_keeps_platform_strategy_separate_from_ai_prose():
    from task_server.api_testing.services.load_report_export_service import report_sections
    data={**REPORT,'ai_diagnosis':{'next_run_strategy':{'source':'平台证据策略','reason':'缺少业务监控','objective':'同压力补充采样','can_prefill':False}}}
    sections=dict(report_sections(data)); rows=dict(sections['AI 诊断与建议'])
    assert rows['推荐依据']=='缺少业务监控'
    assert rows['配置可预填'].startswith('否')
